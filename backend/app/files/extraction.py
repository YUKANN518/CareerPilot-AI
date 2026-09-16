from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.exceptions import AppError
from app.schemas.resume import TextBlock


@dataclass(frozen=True)
class ExtractionOutput:
    raw_text: str
    blocks: list[TextBlock]


class ResumeTextExtractor:
    def extract(self, path: Path, file_format: str) -> ExtractionOutput:
        if file_format == "pdf":
            return self._extract_pdf(path)
        if file_format == "docx":
            return self._extract_docx(path)
        raise AppError("UNSUPPORTED_FILE_TYPE", "不支持该简历格式", 415)

    def _extract_pdf(self, path: Path) -> ExtractionOutput:
        try:
            reader = PdfReader(path)
            blocks: list[TextBlock] = []
            for page_number, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                for block_index, block_text in enumerate(self._split_blocks(text)):
                    blocks.append(
                        TextBlock(
                            source_type="page",
                            text=block_text,
                            page_number=page_number,
                            block_index=block_index,
                        )
                    )
        except (PdfReadError, OSError, ValueError) as exc:
            raise AppError("EXTRACTION_FAILED", "PDF 文本提取失败", 422) from exc

        raw_text = "\n\n".join(block.text for block in blocks)
        if len(raw_text.strip()) < 20:
            raise AppError(
                "SCANNED_PDF_UNSUPPORTED",
                "未检测到可用文本层，暂不支持扫描版 PDF 或需要 OCR 的文件",
                422,
            )
        return ExtractionOutput(raw_text=raw_text, blocks=blocks)

    def _extract_docx(self, path: Path) -> ExtractionOutput:
        try:
            document = Document(str(path))
            blocks: list[TextBlock] = []
            for paragraph_index, paragraph in enumerate(document.paragraphs):
                text = self._clean_text(paragraph.text)
                if text:
                    blocks.append(
                        TextBlock(
                            source_type="paragraph",
                            text=text,
                            paragraph_index=paragraph_index,
                        )
                    )
            for table_index, table in enumerate(document.tables):
                for row_index, row in enumerate(table.rows):
                    text = self._clean_text(" | ".join(cell.text for cell in row.cells))
                    if text:
                        blocks.append(
                            TextBlock(
                                source_type="table",
                                text=text,
                                table_index=table_index,
                                row_index=row_index,
                            )
                        )
        except (PackageNotFoundError, OSError, ValueError) as exc:
            raise AppError("EXTRACTION_FAILED", "DOCX 文本提取失败", 422) from exc

        raw_text = "\n\n".join(block.text for block in blocks)
        if len(raw_text.strip()) < 10:
            raise AppError("NO_EXTRACTABLE_TEXT", "DOCX 中没有可提取文本", 422)
        return ExtractionOutput(raw_text=raw_text, blocks=blocks)

    @classmethod
    def _split_blocks(cls, text: str) -> list[str]:
        return [
            cleaned for part in re.split(r"\n\s*\n|\n", text) if (cleaned := cls._clean_text(part))
        ]

    @staticmethod
    def _clean_text(text: str) -> str:
        return re.sub(r"[ \t]+", " ", text).strip()
