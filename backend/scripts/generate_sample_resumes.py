from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas

OUTPUT_DIRECTORY = Path(__file__).resolve().parents[1] / "tests" / "fixtures"
LINES = [
    "Sample Candidate",
    "sample.candidate@example.test | +86 138 0000 0000 | Shanghai",
    "Summary: Backend engineer building reliable web services.",
    "Work Experience: Software Engineer Intern at Example Labs.",
    "Project: Career planning service using Python, FastAPI and SQLAlchemy.",
    "Skills: Python, FastAPI, SQLAlchemy, SQLite, Vue, TypeScript, Communication.",
    "Education: Example University, Bachelor of Computer Science.",
    "Languages: English and Chinese.",
]


def generate_pdf(path: Path) -> None:
    canvas = Canvas(str(path), pagesize=A4)
    text = canvas.beginText(72, 780)
    text.setLeading(20)
    for line in LINES:
        text.textLine(line)
    canvas.drawText(text)
    canvas.save()


def generate_docx(path: Path) -> None:
    document = Document()
    document.add_heading(LINES[0], level=1)
    for line in LINES[1:]:
        document.add_paragraph(line)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Evidence"
    table.cell(1, 0).text = "Python"
    table.cell(1, 1).text = LINES[4]
    document.save(path)


def main() -> None:
    OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)
    generate_pdf(OUTPUT_DIRECTORY / "sample_resume.pdf")
    generate_docx(OUTPUT_DIRECTORY / "sample_resume.docx")
    print(f"Generated sample resumes in {OUTPUT_DIRECTORY}")


if __name__ == "__main__":
    main()
