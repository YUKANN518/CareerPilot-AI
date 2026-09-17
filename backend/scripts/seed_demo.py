from __future__ import annotations

import argparse
import json
from datetime import UTC, datetime, timedelta
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from sqlalchemy import delete, or_, select

from app import models  # noqa: F401
from app.core.config import get_settings
from app.db.session import SessionLocal
from app.files.storage import PrivateFileStorage
from app.matching.config import DETERMINISTIC_V11_CONFIG
from app.models.enums import JobSourceType, ResumeStatus, SourceStatus, TaskStatus
from app.models.jobs import Job, JobSkill, JobSource
from app.models.matching import MatchReport
from app.models.operations import KnowledgeDocument
from app.models.resumes import FileAsset, Resume, ResumeSkill, ResumeVersion, Skill
from app.models.users import User
from app.services.bootstrap import ensure_local_demo_admin
from app.services.knowledge_documents import KnowledgeDocumentService

DEMO_RESUME_PREFIX = "[DEMO]"
DEMO_SOURCE_PREFIX = "CareerPilot Demo"
DEMO_KNOWLEDGE_PREFIX = "[DEMO]"
SAMPLE_DATA_DIRECTORY = Path(__file__).resolve().parents[2] / "sample_data"
BACKEND_DIRECTORY = Path(__file__).resolve().parents[1]
DEMO_RUNTIME_DIRECTORY = (BACKEND_DIRECTORY.parent / "data" / "runtime" / "demo").resolve()

SKILLS = (
    ("Python", "TECHNICAL"),
    ("FastAPI", "TECHNICAL"),
    ("SQL", "TECHNICAL"),
    ("Vue.js", "TECHNICAL"),
    ("TypeScript", "TECHNICAL"),
)


def evidence(value: str) -> dict[str, object]:
    return {
        "value": value,
        "confidence": 0.96,
        "evidence_text": value,
        "source_location": {
            "source_type": "paragraph",
            "paragraph_index": 1,
            "label": "paragraph:1",
        },
        "needs_confirmation": False,
    }


def empty_evidence() -> dict[str, object]:
    """Represent a schema-required field for which the demo fixture has no evidence."""
    return {
        "value": "",
        "confidence": 0.1,
        "evidence_text": "",
        "source_location": {"source_type": "unknown", "label": "not-found"},
        "needs_confirmation": True,
    }


def load_json_fixture(relative_path: str) -> Any:
    path = SAMPLE_DATA_DIRECTORY / relative_path
    if not path.is_file():
        raise RuntimeError(f"Missing demo fixture: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def profile(persona: dict[str, Any]) -> dict[str, object]:
    education = persona["education"]
    experience = persona["experience"]
    project = persona["project"]
    return {
        "basic_info": {
            "full_name": evidence(persona["full_name"]),
            "email": evidence(persona["email"]),
            "phone": empty_evidence(),
            "location": evidence(persona["location"]),
        },
        "education": [
            {
                "institution": evidence(education["institution"]),
                "degree": evidence(education["degree"]),
                "field_of_study": evidence(education["field_of_study"]),
                "start_date": empty_evidence(),
                "end_date": empty_evidence(),
                "description": empty_evidence(),
            }
        ],
        "work_experience": [
            {
                "company": evidence(experience["company"]),
                "title": evidence(experience["title"]),
                "start_date": empty_evidence(),
                "end_date": empty_evidence(),
                "description": evidence(experience["description"]),
            }
        ],
        "project_experience": [
            {
                "name": evidence(project["name"]),
                "role": empty_evidence(),
                "start_date": empty_evidence(),
                "end_date": empty_evidence(),
                "description": evidence(project["description"]),
            }
        ],
        "technical_skills": [evidence(item) for item in persona["skills"]],
        "soft_skills": [evidence("Cross-functional communication")],
        "languages": [evidence(item) for item in persona["languages"]],
        "certificates": [],
        "awards": [],
        "summary": evidence(persona["summary"]),
    }


def markdown_as_docx(content: str) -> bytes:
    """Convert tracked Markdown fixtures to the DOCX format supported by the KB."""
    document = Document()
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            document.add_heading(line.removeprefix("# "), level=1)
        else:
            document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def assert_isolated_demo_runtime(settings: Any) -> None:
    if settings.app_env.casefold() != "demo":
        raise RuntimeError("Refusing to seed data unless APP_ENV=demo")
    if not settings.database_url.startswith("sqlite:///"):
        raise RuntimeError("Demo seed requires an isolated SQLite database")
    configured_paths = {
        "DATABASE_URL": Path(settings.database_url.removeprefix("sqlite:///")),
        "UPLOAD_DIRECTORY": Path(settings.upload_directory),
        "FAISS_INDEX_DIR": Path(settings.faiss_index_dir),
    }
    for name, configured_path in configured_paths.items():
        resolved = (
            configured_path
            if configured_path.is_absolute()
            else BACKEND_DIRECTORY / configured_path
        ).resolve()
        if not resolved.is_relative_to(DEMO_RUNTIME_DIRECTORY):
            raise RuntimeError(
                f"Refusing demo seed because {name} is outside {DEMO_RUNTIME_DIRECTORY}"
            )


def get_or_create_skill(session, name: str, category: str) -> Skill:
    skill = session.scalar(select(Skill).where(Skill.name == name))
    if skill is None:
        skill = Skill(name=name, category=category)
        session.add(skill)
        session.flush()
    return skill


def seed() -> None:
    settings = get_settings()
    assert_isolated_demo_runtime(settings)
    persona: dict[str, Any] = load_json_fixture("resume/alex_chan.json")
    job_fixtures: list[dict[str, Any]] = load_json_fixture("jobs/jobs.json")
    with SessionLocal() as session:
        ensure_local_demo_admin(session, settings)
        admin = session.scalar(
            select(User).where(User.email == settings.demo_admin_email.strip().casefold())
        )
        if admin is None:
            raise RuntimeError(
                "Demo admin is disabled. Set APP_ENV=demo and DEMO_ADMIN_ENABLED=true."
            )

        skill_map = {
            name: get_or_create_skill(session, name, category) for name, category in SKILLS
        }

        resume_specs = ((persona["title"], persona["skills"]),)
        versions: list[ResumeVersion] = []
        for index, (title, resume_skills) in enumerate(resume_specs, start=1):
            resume_title = f"{DEMO_RESUME_PREFIX} {title}"
            resume = session.scalar(
                select(Resume).where(
                    Resume.owner_id == admin.id,
                    Resume.title == resume_title,
                )
            )
            if resume is None:
                resume = Resume(
                    owner_id=admin.id,
                    title=resume_title,
                    status=ResumeStatus.CONFIRMED,
                    confirmed_at=datetime.now(UTC) - timedelta(days=index),
                )
                session.add(resume)
                session.flush()
                version = ResumeVersion(
                    resume_id=resume.id,
                    version_number=1,
                    raw_text=persona["summary"],
                    structured_data=profile(persona),
                    is_current=True,
                    is_confirmed=True,
                )
                session.add(version)
                session.flush()
                for skill_index, skill_name in enumerate(resume_skills, start=1):
                    version.skills.append(
                        ResumeSkill(
                            skill_id=skill_map[skill_name].id,
                            raw_name=skill_name,
                            confidence=0.94,
                            evidence_text=f"Used {skill_name} in a fictional portfolio project.",
                            evidence_section="project_experience",
                            evidence_source_id=f"demo-paragraph-{skill_index}",
                            source_location=json.dumps(
                                {
                                    "source_type": "paragraph",
                                    "paragraph_index": skill_index,
                                    "label": f"paragraph:{skill_index}",
                                },
                                ensure_ascii=False,
                            ),
                            is_user_confirmed=True,
                        )
                    )
            current_version = session.scalar(
                select(ResumeVersion).where(
                    ResumeVersion.resume_id == resume.id,
                    ResumeVersion.is_current.is_(True),
                )
            )
            if current_version is not None:
                versions.append(current_version)

        sources: list[JobSource] = []
        for source_type, suffix in (
            (JobSourceType.MANUAL, "精选岗位"),
            (JobSourceType.CSV, "课程样例导入"),
        ):
            name = f"{DEMO_SOURCE_PREFIX} · {suffix}"
            source = session.scalar(select(JobSource).where(JobSource.name == name))
            if source is None:
                source = JobSource(
                    created_by_id=admin.id,
                    name=name,
                    source_type=source_type,
                    status=SourceStatus.ENABLED,
                    base_url="https://careers.example.test",
                    region="HK",
                    language="en-HK",
                    sync_frequency_minutes=1440,
                    is_enabled=True,
                    last_tested_at=datetime.now(UTC),
                    last_test_succeeded=True,
                    last_sync_at=datetime.now(UTC),
                )
                session.add(source)
                session.flush()
            sources.append(source)

        jobs: list[Job] = []
        now = datetime.now(UTC)
        for index, fixture in enumerate(job_fixtures):
            source = sources[index % len(sources)]
            external_id = fixture["external_id"]
            job = session.scalar(
                select(Job).where(
                    Job.source_id == source.id,
                    Job.external_job_id == external_id,
                )
            )
            if job is None:
                content = "|".join(
                    str(fixture[key])
                    for key in ("title", "company", "location", "requirements", "external_id")
                )
                job = Job(
                    source_id=source.id,
                    external_job_id=external_id,
                    title=fixture["title"],
                    company=fixture["company"],
                    location=fixture["location"],
                    salary_min=fixture["salary_min"],
                    salary_max=fixture["salary_max"],
                    currency=fixture["currency"],
                    employment_type=fixture["employment_type"],
                    experience_level=fixture["experience_level"],
                    education_requirement="Bachelor degree or equivalent experience",
                    language_requirements=["English"],
                    description=fixture["description"],
                    responsibilities=fixture["responsibilities"],
                    requirements=fixture["requirements"],
                    source_url=f"https://careers.example.test/jobs/{external_id}",
                    normalized_source_url=f"https://careers.example.test/jobs/{external_id}",
                    published_at=now - timedelta(days=index % 14),
                    content_hash=sha256(content.encode()).hexdigest(),
                    raw_data={"demo_seed": True, "completeness": 0.92},
                    status="ACTIVE",
                )
                session.add(job)
                session.flush()
                for skill_name in fixture["skills"]:
                    skill = skill_map[skill_name]
                    job.skills.append(
                        JobSkill(
                            skill_id=skill.id,
                            is_required=True,
                            weight=2,
                            evidence_text=fixture["requirements"],
                        )
                    )
            jobs.append(job)

        if versions:
            scores = (91.0, 78.0, 94.0, 87.0)
            for index, score in enumerate(scores):
                job = jobs[index]
                existing = session.scalar(
                    select(MatchReport).where(
                        MatchReport.user_id == admin.id,
                        MatchReport.resume_version_id == versions[0].id,
                        MatchReport.job_id == job.id,
                        MatchReport.scoring_version == "deterministic-v1.1",
                    )
                )
                if existing is None:
                    session.add(
                        MatchReport(
                            user_id=admin.id,
                            resume_version_id=versions[0].id,
                            job_id=job.id,
                            status=TaskStatus.SUCCEEDED,
                            final_score=score,
                            rule_score=score,
                            scoring_version="deterministic-v1.1",
                            scoring_config_snapshot=DETERMINISTIC_V11_CONFIG.model_dump(
                                mode="json"
                            ),
                            recommendation_level=(
                                "STRONGLY_RECOMMENDED"
                                if score >= 90
                                else "RECOMMENDED"
                                if score >= 75
                                else "CONSIDER"
                            ),
                            explanation="Synthetic demo report generated from anonymous fixtures.",
                        )
                    )

        session.commit()

        storage = PrivateFileStorage(settings.upload_directory)
        # The demo resume is synthetic, but the resume API still expects a private
        # source asset so the normal Resume page can render it. Create that asset
        # once in the isolated demo runtime; no user-uploaded file is involved.
        demo_resumes = session.scalars(
            select(Resume).where(
                Resume.owner_id == admin.id,
                Resume.title.startswith(DEMO_RESUME_PREFIX),
            )
        )
        for demo_resume in demo_resumes:
            if demo_resume.file_asset_id is None:
                content = markdown_as_docx(
                    f"# {demo_resume.title}\n\n{persona['summary']}\n\n"
                    f"Skills: {', '.join(persona['skills'])}"
                )
                storage_key = storage.save(admin.id, "docx", content)
                asset = FileAsset(
                    owner_id=admin.id,
                    storage_key=storage_key,
                    original_name="careerpilot-demo-resume.docx",
                    mime_type=(
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    ),
                    size_bytes=len(content),
                    sha256=sha256(content).hexdigest(),
                    file_format="docx",
                )
                session.add(asset)
                session.flush()
                demo_resume.file_asset_id = asset.id
        session.commit()

        knowledge_service = KnowledgeDocumentService(session, settings)
        knowledge_count = 0
        for path in sorted((SAMPLE_DATA_DIRECTORY / "knowledge_base").glob("*.md")):
            title = f"{DEMO_KNOWLEDGE_PREFIX} {path.stem.replace('_', ' ').title()}"
            document = session.scalar(
                select(KnowledgeDocument).where(KnowledgeDocument.title == title)
            )
            if document is None:
                content = markdown_as_docx(path.read_text(encoding="utf-8"))
                storage_key = storage.save(admin.id, "docx", content)
                asset = FileAsset(
                    owner_id=admin.id,
                    storage_key=storage_key,
                    original_name=f"{path.stem}.docx",
                    mime_type=(
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    ),
                    size_bytes=len(content),
                    sha256=sha256(content).hexdigest(),
                    file_format="docx",
                )
                session.add(asset)
                session.flush()
                document = KnowledgeDocument(
                    uploaded_by_id=admin.id,
                    file_asset_id=asset.id,
                    title=title,
                    category="demo-guidance",
                    status="PENDING",
                    chunk_count=0,
                    metadata_json={"demo_seed": True, "source_fixture": path.name},
                )
                session.add(document)
                session.commit()
            if document.status not in {"PENDING", "PROCESSING", "READY"}:
                document.status = "PENDING"
                document.error_code = None
                document.error_message = None
                session.commit()
            if document.status != "READY":
                knowledge_service.process_index_job(document.id)
            knowledge_count += 1

        print(
            f"Demo data ready: {len(resume_specs)} resumes, {len(jobs)} jobs, "
            f"{knowledge_count} knowledge documents, {len(sources)} sources."
        )


def reset() -> None:
    settings = get_settings()
    assert_isolated_demo_runtime(settings)
    with SessionLocal() as session:
        ensure_local_demo_admin(session, settings)
        admin = session.scalar(
            select(User).where(User.email == settings.demo_admin_email.strip().casefold())
        )
        if admin is None:
            return
        knowledge_service = KnowledgeDocumentService(session, settings)
        demo_documents = list(
            session.scalars(
                select(KnowledgeDocument).where(
                    KnowledgeDocument.title.startswith(f"{DEMO_KNOWLEDGE_PREFIX} ")
                )
            )
        )
        for document in demo_documents:
            knowledge_service.delete(document.id)
        demo_resumes = list(
            session.scalars(
                select(Resume).where(
                    Resume.owner_id == admin.id,
                    Resume.title.startswith(DEMO_RESUME_PREFIX),
                )
            )
        )
        demo_resume_ids = [resume.id for resume in demo_resumes]
        demo_version_ids = list(
            session.scalars(
                select(ResumeVersion.id).where(ResumeVersion.resume_id.in_(demo_resume_ids))
            )
        )
        demo_job_ids = list(
            session.scalars(select(Job.id).where(Job.raw_data["demo_seed"].as_boolean().is_(True)))
        )
        if demo_version_ids or demo_job_ids:
            conditions = []
            if demo_version_ids:
                conditions.append(MatchReport.resume_version_id.in_(demo_version_ids))
            if demo_job_ids:
                conditions.append(MatchReport.job_id.in_(demo_job_ids))
            session.execute(delete(MatchReport).where(or_(*conditions)))
        if demo_resume_ids:
            session.execute(delete(Resume).where(Resume.id.in_(demo_resume_ids)))
        if demo_job_ids:
            session.execute(delete(Job).where(Job.id.in_(demo_job_ids)))
        session.execute(delete(JobSource).where(JobSource.name.startswith(DEMO_SOURCE_PREFIX)))
        session.commit()
        print("Demo data removed. The local demo administrator was preserved.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Seed or reset anonymous demo data.")
    parser.add_argument("--reset", action="store_true")
    parser.add_argument("--reset-and-seed", action="store_true")
    args = parser.parse_args()
    if args.reset or args.reset_and_seed:
        reset()
    if not args.reset or args.reset_and_seed:
        seed()


if __name__ == "__main__":
    main()
