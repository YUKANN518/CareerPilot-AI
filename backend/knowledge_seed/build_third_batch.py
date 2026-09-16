"""Build the third, role-focused CareerPilot knowledge-document batch.

Sources are intentionally placed after the topic content.  This preserves the
keywords-first retrieval structure that was validated in the earlier batches.
"""

# ruff: noqa: E501

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).parent
TODAY = "2026-07-24"
LABOUR_GUIDE = "https://www.labour.gov.hk/eng/public/eip/JobHunting_Eng.pdf"
IANG_URL = "https://www.immd.gov.hk/eng/services/visas/IANG.html"


TOPICS: list[dict[str, object]] = [
    {
        "title": "软件开发岗位介绍与准备",
        "category": "it_roles",
        "nature": "通用职业建议",
        "summary": "从交付可维护的软件功能、理解需求到用可运行项目证明开发能力，帮助初级求职者有针对性地准备。",
        "keywords": "软件开发、后端、前端、接口、代码评审、单元测试、Git、可运行项目、技术取舍",
        "sources": [LABOUR_GUIDE, "https://www.onetonline.org/link/summary/15-1252.00"],
        "source_types": ["香港劳工处公开求职指南", "O*NET OnLine 公开职业资料"],
        "sections": [
            ("一、这类岗位解决什么问题", "初级软件开发岗位的重点不是背出某一种框架，而是把需求转成可运行、可维护、可验证的功能。日常工作可能包括阅读已有代码、拆分小任务、实现页面或接口、处理异常、写测试、记录变更，并和产品、设计、测试或其他开发者协作。不同团队的技术栈和职责范围会不同，因此职位描述中的语言、平台、交付物和协作方式比岗位名称更值得逐项核对。"),
            ("二、常见能力要求", "可从四个方向准备：第一，能用一门主力语言完成基础逻辑并解释代码；第二，理解 HTTP、数据存取、输入校验和错误处理等常见概念；第三，知道如何用 Git 管理分支和提交；第四，能说明测试、调试和代码评审为何有用。遇到未接触过的工具时，诚实说明已有相近经验和学习方法，比把关键词全部写成“熟练”更可靠。"),
            ("三、适合的人群与入门路径", "喜欢把模糊问题拆成步骤、愿意反复调试并能耐心阅读资料的人，通常更能适应开发工作。初学者可先选一个小而完整的主题，例如校园活动报名、个人记账或图书借阅：先画出用户流程，再实现一个可操作的最小版本。每完成一项功能，就记录输入、输出、异常情况和自己做过的取舍，而不是不断添加没有验证过的功能。"),
            ("四、怎样准备可追问的项目证据", "一个适合展示的 API 项目可以包含登录以外的公开查询接口、创建或更新数据的校验规则、清晰的错误返回，以及几条自动化测试。简历中可写“为活动报名服务设计创建与查询接口，为重复提交增加校验，并为异常输入补充测试”；面试时要能说明数据模型、接口边界、如何复现问题及后来怎样改进。代码仓库、简短 README、运行步骤和测试结果都属于证据的一部分。"),
            ("五、避免把开发准备做成关键词清单", "不要把课程作业直接改成一长串技术名词，也不要承诺从未做过的高并发、部署或安全能力。较好的做法是把项目中的个人模块、使用的工具、验证方式和限制写清楚。例如，若项目只是本地演示，就说明它尚未处理多人并发；若由小组完成，也说明自己负责的接口或页面。真实边界能让后续追问有依据。"),
        ],
    },
    {
        "title": "数据分析岗位介绍与准备",
        "category": "it_roles",
        "nature": "通用职业建议",
        "summary": "区分数据整理、分析解释与业务沟通，指导初级求职者用可复现的数据项目展示能力。",
        "keywords": "数据分析、数据清洗、SQL、数据验证、可视化、仪表盘、分析结论、口径、可复现",
        "sources": [LABOUR_GUIDE, "https://www.onetonline.org/link/summary/15-2051.01"],
        "source_types": ["香港劳工处公开求职指南", "O*NET OnLine 公开职业资料"],
        "sections": [
            ("一、岗位的核心是把数据变成可解释的信息", "初级数据分析工作常围绕收集、核对、整理和解释数据展开，例如使用 SQL 查询数据、检查缺失值和重复值、制作图表或仪表盘、把发现写成能让非技术同事理解的说明。工作并不等于只会画图：每个数字都需要明确来源、时间范围和计算口径。职位名称可能是 Data Analyst、BI Analyst 或业务分析相关岗位，实际任务应以招聘描述为准。"),
            ("二、常见能力要求", "基础准备可包括表格工具、SQL、一个用于清洗或分析的脚本环境，以及可视化工具。更重要的是能解释：数据从哪里来，哪些记录被排除，指标如何计算，图表能支持什么结论、又不能支持什么结论。分析报告中应把事实、假设和待确认问题分开。这样的表达比只列出 Python、Excel 或 Tableau 更能说明工作方式。"),
            ("三、适合的人群与入门路径", "喜欢从杂乱信息中找规律、愿意检查细节并能把结论说清楚的人，可以考虑数据分析方向。入门时不必先寻找“大数据”题目，可使用公开且允许使用的小型数据集，完成“问题—数据检查—清洗—查询—可视化—结论—限制”这一闭环。先把一个指标定义写清楚，再做图表，能避免把图形效果误当成分析质量。"),
            ("四、怎样用项目证明能力", "可准备一个销售记录、校园问卷或服务请求数据项目：在 README 中说明字段含义和数据来源；用 SQL 展示筛选、分组或关联；记录如何处理空值、异常值与重复记录；最后用两三张图回答具体问题。简历可以写“清洗含重复与缺失字段的记录，编写 SQL 汇总规则，并在仪表盘旁标注指标口径与限制”。不要把相关性写成因果，也不要从样本推出未被数据支持的结论。"),
            ("五、与软件开发岗位的不同", "软件开发通常以实现和维护功能为主要交付，数据分析则以可信的查询、解释和沟通为主要交付。两者都需要代码和协作，但数据分析项目要特别保留数据字典、查询逻辑和结论限制。若你更享受追踪一项功能为何失败，可尝试开发或测试；若你更关心数据如何支持判断，并愿意反复核对口径，数据分析会更贴近你的兴趣。"),
        ],
    },
    {
        "title": "软件测试与 QA 岗位介绍",
        "category": "it_roles",
        "nature": "通用职业建议",
        "summary": "说明 QA 如何通过测试设计、缺陷沟通和复测降低交付风险，并给出初学者的项目证据方向。",
        "keywords": "软件测试、QA、测试场景、测试用例、缺陷报告、复测、回归测试、边界条件、可复现步骤",
        "sources": [LABOUR_GUIDE, "https://www.onetonline.org/link/summary/15-1253.00"],
        "source_types": ["香港劳工处公开求职指南", "O*NET OnLine 公开职业资料"],
        "sections": [
            ("一、QA 的工作重点", "软件测试与 QA 的目标是尽早发现、描述和降低产品使用中的风险，而不只是“点一遍页面”。初级岗位可能参与阅读需求、设计测试场景、执行手动测试、记录缺陷、验证修复、进行回归测试，并与开发和产品同事确认预期行为。不同团队会采用不同的流程或自动化工具，但清楚、可复现的沟通始终是核心。"),
            ("二、常见能力要求", "应能从正常流程、异常输入、边界条件和权限差异四个角度设计测试。一个有价值的缺陷报告至少说明环境、前置条件、复现步骤、实际结果、预期结果和必要的截图或日志。初学者还可学习基本的接口测试、浏览器开发工具和简单自动化脚本；但若没有实际自动化经验，不应把工具名称当成已掌握能力。"),
            ("三、适合的人群与入门路径", "愿意仔细观察、喜欢追问“如果输入为空会怎样”、能用事实而非责备语言沟通问题的人，通常适合 QA。可从一个熟悉的网页或课程项目开始：先把用户故事拆成可测试的场景，再写测试用例并实际执行。发现问题后先确认是否稳定复现，避免把一次网络波动直接定义为产品缺陷。"),
            ("四、怎样用项目证明能力", "可建立一个小型测试作品集：选择一个待测 Web 应用，为登录、表单提交和查询功能写测试场景；至少包含一条边界测试和一条异常测试；以编号记录缺陷、优先级依据、复现步骤和修复后的复测结果。简历可以写“为报名表单设计正常、空值与重复提交场景，提交可复现缺陷报告，并在修复后完成回归验证”。这比“负责测试”更容易被追问和核实。"),
            ("五、与开发和 IT Support 的边界", "开发岗位主要实现功能，QA 主要验证行为是否符合已确认的预期；IT Support 则更常处理已上线环境中的用户问题、设备或权限问题。QA 不是替开发者背锅，也不应只报告“有 bug”。准备时要练习把问题说成可验证的事实：在什么条件下、做了什么、看到了什么、应当看到什么，以及下一步需要谁确认。"),
        ],
    },
    {
        "title": "IT Support 岗位介绍",
        "category": "it_roles",
        "nature": "通用职业建议",
        "summary": "介绍 IT Support 的用户支持、诊断、记录和升级职责，帮助初学者用排障过程而非软件关键词建立证据。",
        "keywords": "IT Support、技术支持、工单、故障排查、用户沟通、设备配置、权限、升级、知识库",
        "sources": [LABOUR_GUIDE, "https://www.onetonline.org/link/summary/15-1232.00"],
        "source_types": ["香港劳工处公开求职指南", "O*NET OnLine 公开职业资料"],
        "sections": [
            ("一、这类岗位的主要工作", "IT Support 面向用户和运行中的系统：可能接收工单、协助安装或配置设备与软件、排查账号、网络、打印或应用访问问题，并把处理过程记录下来。无法在授权范围内解决的问题，需要按团队流程升级给更合适的同事或供应商。岗位并非只需要“会修电脑”，还要求在时间压力下确认事实、保护用户信息并保持可交接的记录。"),
            ("二、常见能力要求", "初级求职者可准备基本的操作系统、账号与权限、网络连通性、常用办公工具和设备连接知识。更关键的是诊断顺序：先确认用户现象与影响范围，再检查最容易验证的条件，记录已尝试的步骤，避免未经授权的高风险操作。与用户沟通时，用对方能理解的语言说明下一步和预计反馈方式，不凭猜测承诺解决时间。"),
            ("三、适合的人群与入门路径", "愿意帮助他人、能保持耐心并喜欢按线索逐步排查的人，通常适合 IT Support。可从学校社团、实验室或个人设备的常见问题开始练习：把“无法连接 Wi-Fi”拆成账号、设备、网络和服务端等可能环节。练习目标不是炫耀复杂命令，而是形成安全、可说明、可升级的排障过程。"),
            ("四、怎样用项目或证据证明能力", "可以制作一个模拟支持案例集：为三类问题写工单记录，包含用户描述、影响范围、检查顺序、采取的低风险操作、结果和升级条件；再为其中一类问题写一页简短知识库说明。简历可写“整理常见账号访问问题的排查清单，按工单模板记录复现信息与处理结果，并明确无法处理时的升级路径”。不要在案例中暴露真实账号、设备编号或聊天记录。"),
            ("五、与 Business Analyst 的明确区别", "IT Support 的重点是处理用户正在遇到的事件、恢复可用性并留下处置记录；Business Analyst 的重点是理解业务流程、梳理需求并推动各方确认未来要做什么。两者都需要沟通和记录，但 IT Support 的证据更接近工单、排障与知识库，BA 的证据更接近访谈纪要、流程图和需求确认材料。"),
        ],
    },
    {
        "title": "Business Analyst 岗位介绍",
        "category": "it_roles",
        "nature": "通用职业建议",
        "summary": "从需求澄清、流程分析和跨团队确认说明 Business Analyst 的职责，并提供适合初级求职者的证据示例。",
        "keywords": "Business Analyst、业务分析、需求访谈、流程图、用户故事、验收标准、利益相关者、需求确认",
        "sources": [LABOUR_GUIDE, "https://www.onetonline.org/link/summary/13-1111.00"],
        "source_types": ["香港劳工处公开求职指南", "O*NET OnLine 公开职业资料"],
        "sections": [
            ("一、岗位解决的是需求与协作问题", "Business Analyst 常在业务使用者、产品人员与技术团队之间梳理问题：了解现有流程，访谈相关人员，辨别目标、限制与例外情况，把模糊描述整理成可讨论的材料，并跟进确认。不同公司对 BA 的职责划分不相同，有的偏流程优化，有的偏系统需求或数据分析。因此应从职位描述判断它是否要求行业知识、原型、数据分析或项目协调。"),
            ("二、常见能力要求", "初级 BA 应能提出清楚的问题、区分事实与假设、记录访谈结论，并把流程或需求写得让不同角色都能核对。常用表达可以是流程图、用户故事、字段清单、验收条件或会议纪要，具体格式取决于团队。能力不在于背术语，而在于能解释一个需求从哪里来、谁确认、有哪些例外，以及如何判断交付是否符合预期。"),
            ("三、适合的人群与入门路径", "对“为什么要这样做”感兴趣、乐于倾听不同观点并能把讨论整理成结构化内容的人，可以考虑 BA。入门练习可选择校园活动报名或设备借用流程：分别访谈一位使用者和一位管理员，画出现有步骤，标出重复录入或信息缺失处，再提出小范围改进。不要把个人猜测写成用户需求，应把待确认问题明确列出。"),
            ("四、怎样用项目证明能力", "一个可展示的 BA 案例可包含访谈问题、匿名化的访谈纪要、当前与目标流程图、两三条用户故事及可核对的验收条件。例如，“管理员能查看待审核名单；当资料缺失时系统提示补充；审核结果可被申请人查看”。若做了原型，要说明原型用于讨论而非已完成的产品。这样能证明你会把需求变成可确认内容。"),
            ("五、与 IT Support 的区别", "BA 关注未来流程和系统应该满足哪些需要，常通过访谈、流程分析和确认减少返工；IT Support 关注当前用户遇到的事件，常通过排查、工单和升级恢复服务。两者都重视沟通，但 BA 不应越过用户替其决定业务规则，IT Support 也不应在未经确认时把临时处理当作长期需求。选择时可看自己更享受需求澄清，还是即时排障与用户支持。"),
        ],
    },
    {
        "title": "香港 IANG 与非本地毕业生求职提示",
        "category": "hong_kong_job_search",
        "nature": "官方政策整理",
        "summary": "根据香港入境事务处当前 IANG 页面梳理非本地毕业生在求职前应核对的资格、时间点、材料与个案限制。",
        "keywords": "IANG、非本地毕业生、香港入境事务处、recent graduate、non-recent graduate、毕业证明、逗留申请、资格核对",
        "sources": [IANG_URL],
        "source_types": ["香港入境事务处 IANG 官方页面"],
        "sections": [
            ("一、先把资料性质说清楚", "本资料只把香港入境事务处 IANG 官方页面的公开信息整理成求职前核对提示，不构成签证、入境或法律意见，也不保证任何个案结果。IANG 页面说明其面向符合定义的非本地毕业生及相关 GBA 校园毕业生；具体资格、适用范围和个案决定均以申请当时入境事务处公布的信息与处理结果为准。求职时应避免把“正在了解 IANG”写成已经获批或可以免除所有手续。"),
            ("二、理解 recent 与 non-recent 的时间点", "官方页面把在毕业证书所示毕业日期后六个月内递交申请的非本地毕业生列为 recent graduate，超过六个月递交的列为 non-recent graduate。两类申请安排并不完全相同：官方页面列明 recent graduate 在申请时不一定需要已取得聘用；non-recent graduate 的情况则有不同要求。求职者应先记录自己的毕业日期、拟申请日期和学校资格资料，再阅读官方原页的当前说明。"),
            ("三、求职材料与申请材料要分开管理", "简历、作品集和职位申请表用于说明你能胜任岗位；入境申请材料则应按官方要求准备。官方页面列出有效旅行证件、香港身份证（如有）、成绩单、毕业证书或由院校出具的相关证明等材料类别。不要把隐私文件上传到非官方收集表单，也不要因招聘方口头要求而提交与岗位无关的证件副本。文件是否需要补充、翻译或采用特定格式，应以官方页面和正式通知为准。"),
            ("四、把求职节奏与资格核对并行", "在准备香港初级 IT 岗位时，可以同时维护一份事实清单：学位和课程是否符合官方描述、毕业文件是否可取得、申请时间点、职位申请状态及需要向学校或官方渠道确认的问题。招聘流程、雇主决定和入境申请是不同环节；收到面试或职位信息不等于入境申请已获处理。与雇主沟通时，可如实说明自己现有身份和正在依据官方信息核对的安排，不应作出未经批准的承诺。"),
            ("五、变化与个案限制", "入境事务处页面明确提示资格准则可能随时调整，申请会按个别情况处理。因而不要依赖社交平台的旧经验、他人的材料清单或本资料替代官方页面。递交前应再次查看 IANG 原页、所链接的申请渠道和所需证明；若个人情况复杂，应直接向香港入境事务处或合资格专业人士取得针对个案的意见。本资料刻意不把任何条件写成必然获批或法律承诺。"),
        ],
    },
    {
        "title": "初级 IT 岗位选择方法",
        "category": "career_planning",
        "nature": "CareerPilot 方法指南",
        "summary": "用日常任务、能力证据、学习成本和反馈偏好比较初级 IT 方向，而不以薪资或就业传闻代替个人判断。",
        "keywords": "初级 IT 岗位选择、软件开发、数据分析、QA、IT Support、Business Analyst、能力证据、试做项目",
        "sources": [LABOUR_GUIDE, "https://www.onetonline.org/link/summary/15-1252.00", "https://www.onetonline.org/link/summary/15-2051.01", "https://www.onetonline.org/link/summary/15-1253.00", "https://www.onetonline.org/link/summary/15-1232.00", "https://www.onetonline.org/link/summary/13-1111.00"],
        "source_types": ["香港劳工处公开求职指南", "O*NET OnLine 公开职业资料"],
        "sections": [
            ("一、不要先问哪个方向“最好”", "初级岗位选择应从真实工作内容开始，而不是从岗位名称、传闻或单一技能证书开始。软件开发偏向实现与维护功能；数据分析偏向整理、核对和解释数据；QA 偏向验证预期行为与报告风险；IT Support 偏向响应用户事件和排障；Business Analyst 偏向梳理需求与流程。每个岗位都可能需要沟通和基础技术能力，但交付物与反馈方式明显不同。"),
            ("二、用四个维度比较", "第一，看自己愿意长期练习的任务：写与调试代码、核对数据、设计测试、排查用户问题，还是访谈并整理流程。第二，看现有证据：是否已有可运行功能、SQL 分析、测试记录、工单式案例或需求材料。第三，看当前差距能否在一个小项目中补上。第四，看自己接受反馈的方式：开发看功能和代码，QA 看复现与覆盖，Support 看处置记录，BA 看相关方确认，数据分析看口径和解释。"),
            ("三、用短试做替代空想", "给每个感兴趣方向安排一次小试做：开发完成一个含校验的接口；数据分析清洗一份小数据并写出结论限制；QA 为同一功能设计边界场景并提交缺陷报告；IT Support 写一条排障与升级工单；BA 访谈一位模拟使用者并画出流程。试做后记录自己卡在哪里、是否愿意继续迭代、需要什么反馈。试做不能替代真实经验，却能帮助发现偏好与待补能力。"),
            ("四、把证据与职位描述对应", "阅读目标职位时，把要求分成“已有证据”“可在短期补充”“暂时没有”。软件开发可展示接口、代码说明和测试；数据分析可展示清洗、SQL、图表及限制；QA 可展示测试场景、缺陷与复测；IT Support 可展示工单、排查与知识库；BA 可展示访谈、流程图和验收条件。不要因为某个方向缺一项工具就否定自己，也不要把别人的项目证据写成自己的。"),
            ("五、形成可调整的选择结论", "选择一个主要方向和一个相邻方向即可，例如先投 QA，同时保留 IT Support；或先准备数据分析，同时补充 BA 的需求表达。每两周用新项目、面试反馈或职位描述回看证据清单，而不是频繁更换所有学习内容。若发现某岗位的日常任务与自己明显不匹配，可以调整方向；调整意味着更新证据与计划，不等于掩盖过往经验或承诺不存在的能力。"),
        ],
    },
    {
        "title": "入职前沟通与职场基本习惯",
        "category": "workplace_basics",
        "nature": "通用求职建议",
        "summary": "帮助首次入职的求职者以正式渠道确认信息、建立清楚沟通和记录习惯，并保护个人资料。",
        "keywords": "入职前沟通、正式渠道、开始日期、入职文件、工作邮件、沟通记录、信息安全、职场习惯、确认事项",
        "sources": [LABOUR_GUIDE, "https://seng.hkust.edu.hk/iei"],
        "source_types": ["香港劳工处公开求职指南", "大学职业发展公开资源"],
        "sections": [
            ("一、以正式渠道确认关键事实", "收到口头通知或电子信息后，先通过公司公开网站、正式招聘邮箱或已核实的联系人确认职位名称、团队、开始日期、地点或远程安排、报到时间、需要携带或提交的文件，以及下一位联络人。把尚未确认的事项整理成简短问题一次询问，避免在多个聊天窗口反复追问。未收到书面确认前，不应凭个人推测安排不可撤销的搬迁、离职或费用支出。"),
            ("二、写一封清楚而克制的确认邮件", "可采用“感谢—复述事实—列出待确认问题—说明下一步”的结构。例如：感谢对方说明安排；复述已知的开始日期和岗位；询问报到地点、首日联系人及文件提交方式；最后说明会在收到指示后完成准备。邮件不需要过长，也不应夹带与当前入职无关的大量个人经历。发送前检查收件人、附件和敏感信息，避免把身份证件或银行资料误发给未核实的地址。"),
            ("三、首周的基本工作习惯", "新成员可以在第一周建立自己的任务记录：每项任务写下目标、负责人、截止时间、可交付物和不确定处。遇到不明白的缩写、权限或优先级，先查看已有文档，再在合适时间带着具体问题询问。完成任务后，用一两句总结已做内容、结果和下一步，便于同事接手或反馈。记录不是监控同事，而是减少遗漏和误解。"),
            ("四、处理不清楚或有风险的要求", "当指示不完整时，可先复述理解：“我理解目标是……，我计划先……，请确认优先级是否正确。”如果任务涉及生产数据、用户资料、付款、权限或对外发布，应先确认授权范围和流程，不要为了显得积极而绕过审批或使用私人账号传输文件。遇到与正式联系人信息不一致的请求，应暂停并通过已核实渠道再次确认。"),
            ("五、建立可靠的职业形象", "可靠并不等于随时在线，而是对承诺、进度和风险有清楚反馈。无法按时完成时，尽早说明目前进展、阻碍和需要的协助；收到修改意见时，先确认事实和下一步，而非立即辩解。保护公司、用户和同事信息，不在公开作品集、群聊或个人设备中复制敏感内容。良好习惯会在每一次小交付中被看见，也需要持续学习团队自己的规范。"),
        ],
    },
]


def configure_document(document: Document) -> None:
    section = document.sections[0]
    for attribute in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attribute, Inches(1))
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    for name, size, before, after, color in [
        ("Heading 1", 16, 18, 10, "2E74B5"),
        ("Heading 2", 13, 14, 7, "2E74B5"),
        ("Heading 3", 12, 10, 5, "1F4D78"),
    ]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_paragraph(document: Document, text: str, *, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)


def write_document(topic: dict[str, object], path: Path) -> None:
    document = Document()
    configure_document(document)
    title = str(topic["title"])
    title_paragraph = document.add_paragraph()
    title_paragraph.paragraph_format.space_after = Pt(8)
    title_run = title_paragraph.add_run(title)
    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(11, 37, 69)
    metadata = [
        f"标题：{title}",
        f"分类：{topic['category']}",
        "目标读者：香港大学生、应届毕业生和初级 IT 求职者",
        "适用地区：香港",
        f"更新时间：{TODAY}",
        f"资料性质：{topic['nature']}",
        f"检索关键词：{topic['keywords']}",
        f"资料定位：{topic['summary']}",
    ]
    for item in metadata:
        add_paragraph(document, item)
    for heading, body in topic["sections"]:  # type: ignore[index]
        add_paragraph(document, heading, style="Heading 1")
        add_paragraph(document, body)
    add_paragraph(document, "来源说明", style="Heading 1")
    add_paragraph(document, "来源列表：" + "；".join(topic["sources"]))  # type: ignore[arg-type,index]
    add_paragraph(
        document,
        "本资料以公开资料为基础进行中文整理与改写，不复制来源原文，不提供薪资、就业率、录用结果或个案法律结论。岗位职责会随雇主、团队和职位描述而变；涉及 IANG 或逗留条件时，请以香港入境事务处申请当日的正式说明为准。",
    )
    document.save(path)


def main() -> None:
    manifest_path = ROOT / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else []
    batch_titles = {str(topic["title"]) for topic in TOPICS}
    retained = [item for item in manifest if item["title"] not in batch_titles]
    existing = {item["title"] for item in retained}
    additions = []
    for topic in TOPICS:
        title = str(topic["title"])
        if title in existing:
            raise ValueError(f"Duplicate knowledge title: {title}")
        category = str(topic["category"])
        filename = f"{category}/{title}.docx"
        path = ROOT / filename
        path.parent.mkdir(parents=True, exist_ok=True)
        write_document(topic, path)
        additions.append(
            {
                "filename": filename,
                "title": title,
                "category": category,
                "language": "zh-CN",
                "region": "Hong Kong",
                "updated_at": TODAY,
                "source_urls": topic["sources"],
                "source_types": topic["source_types"],
                "content_summary": topic["summary"],
            }
        )
    manifest_path.write_text(json.dumps([*retained, *additions], ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"created": additions, "total_manifest": len(retained) + len(additions)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
