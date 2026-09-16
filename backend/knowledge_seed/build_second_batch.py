from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

# ruff: noqa: E501

ROOT = Path(__file__).parent
TODAY = "2026-07-24"
LABOUR_GUIDE = "https://www.labour.gov.hk/eng/public/eip/JobHunting_Eng.pdf"
YES_JOB_SEEKING = "https://www.yes.labour.gov.hk/Applicants/Main/JobSeeking"
LABOUR_SUMMER = "https://www.labour.gov.hk/text_alternative/pdf/eng/Guide_for_Students_Seeking_Summer_Jobs.pdf"
LABOUR_TRAPS = "https://www.labour.gov.hk/eng/public/eip/BewareEmployTrap_Eng.pdf"
POLICE_SCAMETER = "https://www.police.gov.hk/offbeat/1219/eng/9027.html"
POLICE_JOB_SCAMS = "https://www.police.gov.hk/offbeat120/scam/49_common-job-scams.html"
HKUST_CAREER = "https://seng.hkust.edu.hk/iei"

TOPICS = [
    {
        "category": "interview",
        "title": "STAR 行为面试回答",
        "nature": "CareerPilot 方法指南",
        "sources": [LABOUR_GUIDE, YES_JOB_SEEKING],
        "source_types": ["香港劳工处公开求职指南", "香港劳工处青年就业服务"],
        "summary": "把行为面试题转化为可核实的个人行动与学习证据。",
        "sections": [
            ("一、先识别行为题真正想问什么", "“请说一次你解决冲突的经历”“你如何处理时间压力”通常不是要求背诵性格优点，而是希望了解你在具体情境中的判断、沟通和行动。回答前先把题目中的能力词圈出，例如协作、主动性、责任感或学习能力；再从课程项目、社团、实习、志愿服务或真实自学经历中选择最贴近的一件事。没有相关经历时，不应虚构职场故事，可以诚实说明自己会如何处理，并补充一个相邻的真实例子。"),
            ("二、用 STAR 建立一组可复用故事", "Situation 用一两句交代背景和限制；Task 说明你个人承担的目标；Action 重点写自己做了哪些决策、沟通或技术操作；Result 只写能解释或展示的结果。准备时不必为每道题背一篇稿，而可建立四到六个故事库，例如一次修复缺陷、一次协调分工、一次学习新工具和一次面对失败。每个故事都标注可回答的能力词，面试时再按问题调整重点。"),
            ("三、案例：把“团队合作很好”变成证据", "课程小组在演示前发现报名页面会重复写入资料。背景是截止时间临近，任务是让重复提交有明确处理。行动可以说：我先复现问题，与负责前端的同学确认触发路径，再在接口加入唯一性校验并补充两条测试。结果是演示时重复提交会返回提示，管理员名单不再产生重复记录。这个案例没有声称提升了某个百分比，却能说明协作、排查和验证。"),
            ("四、回答后的自检", "完成后检查四件事：是否说清个人而非团队总贡献；行动是否具体到方法或沟通对象；结果是否可验证；是否直接回应题目要求的能力。若面试官追问“为什么这样做”或“下次会怎样改进”，可补充当时的限制和复盘，而不是临时添加没有发生的成果。"),
        ],
    },
    {
        "category": "interview",
        "title": "项目深挖问题准备",
        "nature": "CareerPilot 方法指南",
        "sources": [LABOUR_GUIDE, HKUST_CAREER],
        "source_types": ["香港劳工处公开求职指南", "大学职业发展公开资源"],
        "summary": "为项目面试中的架构、取舍、测试和个人贡献追问建立证据地图。",
        "sections": [
            ("一、把项目从展示稿还原为工作过程", "项目深挖常从“这个项目做什么”开始，随后会问为什么选该方案、你负责哪一部分、遇过什么问题以及如何确认结果。只准备功能介绍通常不够。请为每个项目写一页证据地图：目标用户、输入与输出、个人负责的模块、采用的技术、关键限制、测试方式、仍未解决的部分和可展示的材料。地图应以你真正看过的代码、文档或演示为准。"),
            ("二、四类高频追问", "第一类是边界：系统不做什么、谁使用它。第二类是取舍：为何使用某个数据库、接口或资料结构，替代方案有什么代价。第三类是排错：出现异常时如何复现、定位和修复。第四类是质量：怎样验证输入、处理错误、写测试或收集反馈。即使是小型课程项目，也可以从这四类问题整理答案；重点不是术语数量，而是推理过程能否自洽。"),
            ("三、案例：解释一个接口设计", "若你负责活动报名 API，可以先说明它接收活动编号与报名资料，返回成功或可读错误。被问到重复报名时，不要只说“加了判断”，而可说明先以测试资料复现重复写入，再在服务层检查已有记录，并让前端显示错误信息。被问到限制时，可承认课程项目没有处理高并发或权限分级，并说明下一步会补充的测试，而不是把未完成能力包装成生产经验。"),
            ("四、准备材料而非背答案", "每个项目保留一个可打开的仓库、截图、架构草图或测试记录即可，注意不泄露同学资料、账号或密钥。面试前挑一项最相关项目，练习两分钟概述和五分钟深挖。若某项技术只是队友使用过，应明确说明自己的观察与协作边界；诚实的范围说明通常比模糊地声称“全栈负责”更容易获得信任。"),
        ],
    },
    {
        "category": "interview",
        "title": "技术面试准备方法",
        "nature": "通用求职建议",
        "sources": [LABOUR_GUIDE, YES_JOB_SEEKING, HKUST_CAREER],
        "source_types": ["香港劳工处公开求职指南", "香港劳工处青年就业服务", "大学职业发展公开资源"],
        "summary": "按目标岗位而非泛泛刷题，准备初级 IT 技术面试。",
        "sections": [
            ("一、从职位描述倒推准备范围", "技术面试的范围取决于岗位。先把职位描述拆成编程语言、数据处理、测试、部署、协作和业务理解六栏，再标注已有证据与待学习项。初级岗位通常更在意你能否解释基础概念、读懂小段代码、说明项目选择和接受反馈，而不是要求你假装掌握所有框架。没有写明的要求不要自行假定，缺失信息应保留为待确认。"),
            ("二、建立“概念—代码—项目”三层复习", "概念层用自己的话解释，例如 HTTP 请求为什么需要状态码、SQL 查询如何处理空值、单元测试在验证什么。代码层练习读题、拆分输入、写小函数和检查边界条件。项目层把概念连接到真实经历，例如在报名项目中怎样验证输入、怎样记录错误。三层能互相支撑：只背定义容易断裂，只背项目又可能回答不了基础追问。"),
            ("三、一次有效练习的节奏", "选一道与岗位相关的小题，先用两分钟复述需求和问清约束，再写出步骤或伪代码，最后检查空输入、重复资料和错误返回。完成后录下解释过程，听自己是否跳过关键假设。若不能完成，也要记录卡在哪里：语法、数据结构、调试还是题意理解。下一次只补一个薄弱点，比在短时间内堆积大量题目更可追踪。"),
            ("四、面试当天的技术表达", "遇到问题时先说明理解，再提出可验证的方案；不确定时可以说“我会先查阅官方文档或写一个最小复现”，并解释为何这样验证。不要运行来历不明的代码，也不要把网上答案当成自己做过的项目。技术准确性重要，但面试官同样能从你的提问、边界意识和复盘方式判断是否适合初级团队协作。"),
        ],
    },
    {
        "category": "interview",
        "title": "面试中如何回答不会的问题",
        "nature": "CareerPilot 方法指南",
        "sources": [LABOUR_GUIDE, YES_JOB_SEEKING],
        "source_types": ["香港劳工处公开求职指南", "香港劳工处青年就业服务"],
        "summary": "在未知技术或经历题面前保持诚实，并展示可验证的学习与排查路径。",
        "sections": [
            ("一、区分“完全不会”与“需要澄清”", "有些题目并非你不会，而是范围不清。例如“你会优化数据库吗”可能涉及索引、查询、资料建模或监控。先复述自己的理解，并问一个窄问题：这里是指课程项目中的查询性能，还是生产环境的监控？澄清能避免答偏，也让面试官看到你不轻易假设。若确实没有经历，直接说明范围比猜测术语安全。"),
            ("二、四步回答框架", "第一步承认边界：“我还没有在生产环境做过这个。”第二步连接相邻经验：“但我在课程项目中处理过重复记录和输入验证。”第三步说明解决路径：“我会先阅读官方文档、建立最小示例、写测试确认行为，再请有经验的同事复核。”第四步回到学习意愿：“如果这个能力是岗位重点，我会把它列为入职前或试用期的学习目标。”这不是承诺已经掌握，而是说明行动方法。"),
            ("三、案例：被问到没有用过的工具", "面试官问你是否用过容器工具时，可以回答：我没有在真实部署环境使用过该工具，因此不会声称有生产经验；我知道它可帮助统一运行环境。在报名系统项目里，我负责接口验证和测试，若要学习这项工具，我会先把项目整理成最小可运行版本，按官方入门资料完成构建，再检查服务能否在另一台机器启动。回答中的“知道”和“做过”被清楚区分。"),
            ("四、避免两种失分做法", "第一种是用“我很快能学会”结束回答，却没有学习路径。第二种是为了显得熟悉而编造经验，后续追问通常会暴露矛盾。可以在回答后请面试官说明团队最常使用的场景，再把问题转化为你愿意研究的方向。诚实并不等于沉默；重点是让未知变成一个可拆解、可验证的下一步。"),
        ],
    },
    {
        "category": "interview",
        "title": "面试结束时如何反问",
        "nature": "通用求职建议",
        "sources": [LABOUR_GUIDE, HKUST_CAREER],
        "source_types": ["香港劳工处公开求职指南", "大学职业发展公开资源"],
        "summary": "用与岗位有关的问题确认工作内容、学习支持与下一步，而不是套用空泛提问。",
        "sections": [
            ("一、反问不是表现聪明的环节", "面试结束时的提问可以帮助你判断岗位是否适合，也让对方看到你认真阅读了职位描述。不要为了提问而提问。先回想面试中仍不清楚的内容：这个初级职位最先交付什么、团队如何协作、谁会提供反馈、下一步流程如何安排。若对方已经解释过，就不要原样重复。"),
            ("二、围绕工作内容提问", "可问“新同事在前三个月通常会参与哪些类型的任务？”“这个岗位与产品、测试或业务同事怎样协作？”“团队如何定义一个功能完成？”这些问题关注实际工作，而非要求对方预测你的录用结果。对于初级 IT 岗位，也可问代码评审、测试或文档习惯，前提是这些内容与岗位描述或面试对话相关。"),
            ("三、围绕成长和评价提问", "Graduate Programme 或实习不一定有统一培养方式，因此可以直接问“是否有导师、培训或轮岗安排，哪些安排已确定，哪些会因团队而异？”还可以问“试用期或阶段性反馈通常关注哪些可观察的工作行为？”回答能帮助你判断学习支持是否具体。不要把未公开的承诺当事实，也不要把对方的概括解释成录用保证。"),
            ("四、围绕流程收尾", "最后可简短确认“后续还有哪些步骤、预计由谁联系我？”如果你需要提供作品集、成绩单或其他资料，也可询问正式提交渠道。薪酬、福利和工作安排属于重要考虑，但是否在首轮询问、怎样询问，应视面试阶段与对方已提供的信息判断。把问题记下并在面试后复盘，下一次可减少重复或过于宽泛的提问。"),
        ],
    },
    {
        "category": "hong_kong_job_search",
        "title": "香港常见招聘流程",
        "nature": "官方政策整理",
        "sources": [LABOUR_GUIDE, YES_JOB_SEEKING, HKUST_CAREER],
        "source_types": ["香港劳工处公开求职指南", "香港劳工处青年就业服务", "大学职业发展公开资源"],
        "summary": "理解香港求职中常见的申请、筛选、面试和后续联络环节，但不把任何一种流程当作保证。",
        "sections": [
            ("一、流程会因雇主和岗位而不同", "香港没有一条适用于所有企业、实习或毕业生岗位的固定招聘路径。公开职位通常会说明申请方式、截止日期、材料和联络渠道；雇主可能先筛选简历，也可能安排测验、电话或网上面试、现场面试、参考资料核实或其他步骤。职位公告没有写出的安排，应标为未提供并通过正式渠道确认，不应从其他公司的经验推断。"),
            ("二、申请阶段要留下可追踪记录", "提交前确认职位名称、版本、附件和申请渠道；提交后记录日期、职位链接、联系人、下一步和自己已提供的材料。劳工处的求职资料提醒求职者阅读职位的职责、资格、要求和申请方式。对学生而言，学校职业中心、招聘讲座、雇主官网和公开职位平台都可能是线索，但公开刊登不等于学校或平台替雇主作担保。"),
            ("三、收到面试或测验邀请后", "先核对邀请来自的域名、联系人、时间、地点或线上链接，再阅读要准备的文件。不同岗位的测验内容不相同：技术岗位可能讨论项目或基础知识，其他岗位可能关注案例、沟通或业务理解。若时间冲突，尽早通过原始邀请中的正式方式沟通；不要把陌生即时通讯帐号发来的链接视为唯一官方渠道。"),
            ("四、结果、跟进与复盘", "没有即时回复不代表某一种固定结果，招聘进度由雇主决定。可以在适当时间用礼貌、简短的邮件询问状态，并继续投递其他合适职位。每次面试后记录被问的问题、回答证据、想补的知识和需要澄清的岗位信息。收到任何录用或入职文件前，先核实公司和联系人，阅读条款；不要因为“限时”压力跳过核对。"),
        ],
    },
    {
        "category": "hong_kong_job_search",
        "title": "香港实习与 Graduate Programme",
        "nature": "通用求职建议",
        "sources": [LABOUR_SUMMER, LABOUR_GUIDE, HKUST_CAREER],
        "source_types": ["香港劳工处学生暑期求职指南", "香港劳工处公开求职指南", "大学职业发展公开资源"],
        "summary": "分辨实习、暑期岗位与 Graduate Programme 的常见目标，按雇主公开资料准备而不假设统一规则。",
        "sections": [
            ("一、先分清机会的性质", "实习可能是课程相关安排、暑期职位或雇主自行设定的短期岗位；Graduate Programme 通常面向毕业生，由雇主设计入职初期的发展安排。名称相同不代表期限、轮岗、培训、资格或转正安排相同。申请前应阅读每个雇主的职位页面、地点、雇佣类型和申请要求；没有写出的内容只记录为待确认，不要把听闻的安排写进求职信。"),
            ("二、用“学习目标”而非光环选择", "对初级 IT 求职者，可比较机会是否让你接触真实任务、获得反馈、理解团队流程，并留下可说明的工作证据。举例说，一个测试实习可能帮助你学习测试场景、缺陷记录和复测；一个开发实习可能让你参与小功能、代码评审或文档整理。不要假设任何实习都会转为全职，也不要承诺自己接受了还未获得的机会。"),
            ("三、准备材料时保留学生身份的真实性", "简历可突出课程项目、社团职责、竞赛或自学产出，但要区分课堂练习、个人作品和真实工作。求职信可说明你希望把哪项能力应用到岗位，却不应把学习计划写成既有经验。若雇主要求成绩单、就读证明或其他资料，应通过公告或正式通知核对版本和提交方式，谨慎处理个人资料。"),
            ("四、实习期间与结束后的复盘", "进入机会后，尽早与主管确认任务边界、反馈方式和可公开展示的成果。每周记录完成内容、学到的工具、收到的反馈和仍有的疑问；这会成为后续面试可以核实的材料。有关学生实习的法定定义或工资安排，必须以劳工处当时公布的适用条件为准；本资料不把某一条规则套用于所有学生岗位。"),
        ],
    },
    {
        "category": "job_application",
        "title": "香港求职诈骗识别",
        "nature": "官方政策整理",
        "sources": [LABOUR_TRAPS, POLICE_SCAMETER, POLICE_JOB_SCAMS, LABOUR_SUMMER],
        "source_types": ["香港劳工处就业陷阱资料", "香港警务处 Scameter 公开资料", "香港警务处防骗资讯", "香港劳工处学生暑期求职指南"],
        "summary": "识别可疑招聘讯号、保护个人资料，并在怀疑受骗时停止操作和寻求官方协助。",
        "sections": [
            ("一、把“异常要求”当作停下来的讯号", "高薪却几乎不要求经验、没有正式面试便催促入职、要求先付款、要求交出银行资料或要求下载不明应用程式，都应触发核实。单一讯号不一定证明诈骗，但越多异常同时出现，越不应在即时通讯工具里仓促决定。劳工处提醒求职者留意以招聘名义骗取金钱或个人资料的就业陷阱。"),
            ("二、三步核实雇主和联络渠道", "第一步从雇主官网、公开商业资料或原职位页面核对公司名称、职位和联络资料。第二步检查邮件网域、网站地址和面试地点是否一致；不要只相信来电者的显示名称。第三步把可疑电话号码、帐号、网址或社交帐号交叉核对。香港警方介绍的 Scameter 可用于查询可疑联络资料的风险线索，但查询结果不是对雇主的录用背书，仍应保留判断。"),
            ("三、个人资料和金钱的底线", "在未核实雇主前，不发送身份证副本、银行帐号、信用卡资料、验证码或网银登入资料；也不要代收或转移不明款项。以培训费、保证金、开户、充值任务或购买指定物品作为入职条件时，应停止付款并核实。真实招聘也可能需要个人资料，但应通过正式、必要且可确认的渠道处理，而不是把所有资料一次传给陌生帐号。"),
            ("四、怀疑受骗后的处理", "保留广告截图、聊天记录、转帐资料、网址和对方帐号，停止继续付款或提供资料，并向可信任的人、学校职业顾问或相关机构求助。香港警方公开资料提到可使用 Scameter 查询，并可致电防骗易热线 18222 寻求意见；紧急或已涉及损失时应按官方指引尽快报警。本资料不替代警方、银行或法律专业人士的个案建议。"),
        ],
    },
]

EXTRAS = {
    "STAR 行为面试回答": (
        "五、把回答控制在可追问的长度",
        "首次回答可控制在约一到两分钟，把最多篇幅留给 Action；若对方感兴趣，再展开技术细节或复盘。准备时可把每个故事录音，删去“非常努力”“效果很好”这类没有证据的词。保留一个改进点也很有用，例如下次会更早确认接口边界或安排测试时间；它说明你能复盘，而不是否定自己的贡献。",
    ),
    "项目深挖问题准备": (
        "五、遇到记不清时怎样处理",
        "项目隔了一段时间后，记不清某个库的版本或具体参数很正常。不要猜一个答案。可以说明自己记得的设计目的，指出会回看哪份提交记录、需求说明或测试文件确认细节。面试前应重新运行项目的关键流程，并把最容易被问到的模块写成简短注释；这样既能恢复记忆，也能避免把网络教程中的做法误说成自己的实现。",
    ),
    "技术面试准备方法": (
        "五、为不同岗位保留不同的复习包",
        "不要把后端、数据分析、测试和 IT Support 的准备混成一份万能清单。后端岗位可优先练接口、资料结构和错误处理；数据岗位可复习数据清洗、查询与结果解释；测试岗位可准备测试设计和缺陷复现；IT Support 岗位则可练问题澄清、记录和升级沟通。每次投递只调整与职位直接相关的部分，避免在简历和面试中造成技能表述不一致。",
    ),
    "面试中如何回答不会的问题": (
        "五、把未知问题变成后续学习记录",
        "面试结束后把不会的问题按“概念不懂、场景不熟、表达不清”分类。概念不懂就找官方文档或可靠课程建立最小示例；场景不熟就把它映射到一个项目实验；表达不清则重新组织成边界、做法和验证。这样下一次遇到同类问题时，答案会来自真实学习产出。不要把一次面试中的题目和答案原样公开为雇主内部资料。",
    ),
    "面试结束时如何反问": (
        "五、根据对象调整问题深度",
        "和未来直属主管交流时，可多问任务边界、协作和反馈；和 HR 或招聘人员交流时，可确认流程、材料和正式联络方式；和团队成员交流时，可了解日常协作体验，但不要要求其代表公司作出决定。准备三到五个问题即可，现场根据已经获得的信息选一两个。若确实没有疑问，礼貌表达感谢并确认下一步，也比提出无关问题更专业。",
    ),
    "香港常见招聘流程": (
        "五、用时间线管理而不是猜测结果",
        "可为每个职位建立一条时间线：看到职位、提交材料、收到通知、参加环节、跟进日期和结果。时间线只记录已发生或已确认的信息，并注明来源，例如雇主邮件、官网公告或学校职业中心通知。它能帮助你发现是否漏交材料，也能在面试前快速回顾岗位要求。不要把别人的等待时间当成自己的承诺，更不要因未回复而向不明帐号补交敏感资料。",
    ),
    "香港实习与 Graduate Programme": (
        "五、申请前的比较提问",
        "比较两个机会时，可写下五个问题：工作地点和出勤安排是什么、最初的任务是什么、谁提供反馈、申请材料与截止日期是什么、哪些信息仍未提供。对于 Graduate Programme，还可问培养安排是否适用于该部门、是否有确定轮岗以及评估方式。问题的作用是补足公开资料，而不是推断录用机会。选择时结合自己的学习目标、时间安排和可承担的通勤或生活条件。",
    ),
    "香港求职诈骗识别": (
        "五、一个安全的处理示例",
        "例如你收到陌生帐号发来的“无需经验、立即开始”兼职邀请，对方要求先下载应用并充值。安全做法不是与其讨价还价，而是先停止点击与付款，保存讯息，再用原公司官网而非对方提供的链接核对职位；如有疑虑，可把电话号码、网站或帐号交给 Scameter 查询，并按警方公开渠道求助。不要为了测试真假而转一笔小额款项，也不要把自己的银行帐户交给对方验证。",
    ),
}

SEARCH_KEYWORDS = {
    "STAR 行为面试回答": "行为面试、Situation、Task、Action、Result、个人贡献、可验证结果",
    "项目深挖问题准备": "项目证据地图、架构取舍、排错、测试、个人模块、项目追问",
    "技术面试准备方法": "技术面试、HTTP 状态码、SQL 空值、单元测试、代码练习、项目复习",
    "面试中如何回答不会的问题": "不会的问题、承认边界、相邻经验、最小示例、学习路径、澄清范围",
    "面试结束时如何反问": "反问、直属主管、任务边界、协作、反馈、导师、轮岗、后续流程",
    "香港常见招聘流程": "香港招聘流程、申请记录、职位链接、联系人、面试邀请、跟进日期、结果",
    "香港实习与 Graduate Programme": "香港实习、暑期职位、Graduate Programme、毕业生、培训、轮岗、导师",
    "香港求职诈骗识别": "求职诈骗、Scameter、充值任务、招聘邀请、邮件网域、个人资料、防骗易 18222",
}


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    for name, size, before, after, color in [
        ("Heading 1", 16, 18, 10, "2E74B5"),
        ("Heading 2", 13, 14, 7, "2E74B5"),
        ("Heading 3", 12, 10, 5, "1F4D78"),
    ]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_paragraph(document: Document, text: str, *, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)


def write_document(topic: dict[str, object], path: Path) -> None:
    document = Document()
    configure_document(document)
    title_paragraph = document.add_paragraph()
    title_paragraph.paragraph_format.space_after = Pt(8)
    title_run = title_paragraph.add_run(str(topic["title"]))
    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(11, 37, 69)
    metadata = [
        f"标题：{topic['title']}",
        f"分类：{topic['category']}",
        "目标读者：香港大学生、应届毕业生和初级 IT 求职者",
        "适用地区：香港",
        f"更新时间：{TODAY}",
        f"资料性质：{topic['nature']}",
    ]
    for item in metadata:
        add_paragraph(document, item)
    add_paragraph(document, "检索关键词：" + SEARCH_KEYWORDS[str(topic["title"])])
    add_paragraph(document, "资料定位：" + str(topic["summary"]))
    for heading, body in topic["sections"]:
        add_paragraph(document, heading, style="Heading 1")
        add_paragraph(document, body)
    extra_heading, extra_body = EXTRAS[str(topic["title"])]
    add_paragraph(document, extra_heading, style="Heading 1")
    add_paragraph(document, extra_body)
    add_paragraph(document, "来源说明", style="Heading 1")
    add_paragraph(document, "来源列表：" + "；".join(topic["sources"]))
    add_paragraph(
        document,
        "本资料为面向学生的中文整理与改写，不复制来源原文，不提供薪资、录用结果或个案法律结论。涉及招聘安排、学生身份或防骗行动时，应以雇主和香港官方机构在当时公布的信息为准。",
    )
    document.save(path)


def main() -> None:
    manifest_path = ROOT / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else []
    batch_titles = {str(topic["title"]) for topic in TOPICS}
    retained_manifest = [item for item in manifest if item["title"] not in batch_titles]
    known_titles = {item["title"] for item in retained_manifest}
    additions = []
    for topic in TOPICS:
        title = str(topic["title"])
        if title in known_titles:
            raise ValueError(f"Duplicate knowledge title: {title}")
        category = str(topic["category"])
        filename = f"{category}/{title}.docx"
        path = ROOT / filename
        path.parent.mkdir(parents=True, exist_ok=True)
        write_document(topic, path)
        additions.append(
            {
                "filename": filename,
                "title": title,
                "category": category,
                "language": "zh-CN",
                "region": "Hong Kong",
                "updated_at": TODAY,
                "source_urls": topic["sources"],
                "source_types": topic["source_types"],
                "content_summary": topic["summary"],
            }
        )
    manifest_path.write_text(
        json.dumps([*retained_manifest, *additions], ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        json.dumps(
            {"created": additions, "total_manifest": len(retained_manifest) + len(additions)},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
